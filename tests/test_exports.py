"""Pruebas de las exportaciones.

El verificación (la suite de contraste) ya compara estas funciones contra el
JavaScript de referencia ejecutando ``exportXLSX``/``exportCSV``/``exportJSONBackup``
con un SheetJS y un ``DB`` falsos. Aquí se fijan los contenidos con valores
calculados a mano (y se comprueba que ``openpyxl`` escribe un libro legible) para
que la suite siga siendo útil sin depender de herramientas externas.

Fechas de referencia: ``today = 2026-03-15T12:30:45`` (hora local, huso UTC−4),
es decir ``1773592245000``.
"""

from __future__ import annotations

import json

import pytest

from app.core.js_locale import (
    iso_string,
    locale_date,
    locale_datetime,
    parse_date_ms,
)
from app.exports import (
    csv_filename,
    backup_filename,
    backup_text,
    build_backup,
    build_closure_act_html,
    build_sheets,
    catalog_rows,
    category_payment_rows,
    column_widths,
    import_plan,
    import_prompt,
    is_valid_backup,
    ledger_rows,
    plan_excel_import,
    products_rows,
    rates_rows as csv_rates_rows,
    to_csv,
    transaction_headers,
    transactions_rows,
    write_xlsx,
    xlsx_filename,
)
from app.exports.excel import analysis_rows, rates_rows, summary_rows

NOW_MS = 1773592245000
NOW_DATE = "2026-03-15"

CATEGORIES = [
    {
        "id": "c1",
        "name": "Comida",
        "color": "#111111",
        "customFields": [
            {"id": "f1", "label": "Marca propia", "type": "boolean"},
            {"id": "f2", "label": "Lote", "type": "text"},
        ],
    },
    {"id": "c2", "name": "Ocio", "color": "#222222"},
]

PRODUCTS = [
    {
        "id": "p1",
        "name": "Harina",
        "categoryId": "c1",
        "brand": "Marca X",
        "unit": "kg",
        "subcategory": "Lácteos",
        "provider": "A",
        "createdAt": 1767240306000,
    },
    {"id": "p2", "name": "Arroz", "categoryId": "c1", "unit": "kg", "createdAt": 1767240306000},
]

RATES = [
    {"id": "r2", "date": "2026-03-01", "value": 60, "source": "BCV"},
    {"id": "r1", "date": "2026-01-01", "value": 40, "source": "BCV"},
]

TXS = [
    {
        "id": 2,
        "workspace": "personal",
        "year": 2026,
        "date": "2026-03-05",
        "productId": "p1",
        "productName": "Harina",
        "categoryId": "c1",
        "subcategory": "Lácteos",
        "brand": "Marca X",
        "provider": "A",
        "taxMode": "incluido",
        "baseImponible": 100,
        "ivaMonto": 16,
        "amountBs": 116,
        "amountUsd": 4,
        "rate": 29,
        "quantity": 2,
        "unit": "kg",
        "notes": 'nota, con "comillas"',
        "customValues": {"f1": True, "f2": "L-9"},
    },
    {
        "id": 1,
        "workspace": "personal",
        "year": 2026,
        "date": "2026-01-10",
        "productId": "p1",
        "productName": "Harina",
        "categoryId": "c1",
        "quantity": None,
        "unit": "kg",
        "amountBs": 80,
        "amountUsd": 2,
        "taxMode": "exento",
        "customValues": {},
    },
    {
        "id": 3,
        "workspace": "personal",
        "year": 2026,
        "date": "2026-03-06",
        "productId": "p2",
        "productName": "Arroz",
        "categoryId": "zz",
        "quantity": 3,
        "unit": "kg",
        "amountBs": 200,
        "amountUsd": 5,
        "taxMode": "incluido",
        "baseImponible": 172.41,
        "ivaMonto": 27.59,
        "customValues": {},
    },
]


# ---------------------------------------------------------------------------
# Fechas con la semántica de JavaScript
# ---------------------------------------------------------------------------


def test_locale_es_ve() -> None:
    assert locale_datetime(NOW_MS) == "15/3/2026, 12:30:45 p. m."
    assert locale_date(NOW_MS) == "15/3/2026"
    # medianoche: la hora 12 con sufijo de mañana
    assert locale_datetime(parse_date_ms("2026-01-01T04:05:06Z")) == (
        "1/1/2026, 12:05:06 a. m."
    )


def test_iso_y_parse() -> None:
    assert iso_string(NOW_MS) == "2026-03-15T16:30:45.000Z"
    assert parse_date_ms("2026-03-15") == 1773532800000.0
    assert parse_date_ms(NOW_MS) == NOW_MS
    assert parse_date_ms("2026-03-15T16:30:45Z") == NOW_MS
    # sin desplazamiento se interpreta en hora local
    assert parse_date_ms("2026-03-15T12:30:45") == NOW_MS


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def test_to_csv_escapa_solo_cuando_hace_falta() -> None:
    text = to_csv(
        [{"a": "x,y", "b": "simple"}, {"a": "sin coma", "b": 'dijo "hola"'}],
        ["a", "b"],
    )
    assert text == 'a,b\n"x,y",simple\nsin coma,"dijo ""hola"""'
    assert to_csv([], ["a", "b"]) == "a,b\n"
    # los saltos de línea obligan a entrecomillar
    assert to_csv([{"a": "linea1\nlinea2"}], ["a"]) == 'a\n"linea1\nlinea2"'


def test_transactions_csv_incluye_campos_personalizados() -> None:
    rows = transactions_rows(TXS, CATEGORIES)
    headers = transaction_headers(rows)
    assert headers[:4] == ["id", "workspace", "fecha", "producto"]
    assert headers[-2:] == ["Var_Marca propia", "Var_Lote"]
    assert rows[0]["Var_Marca propia"] == "Sí"
    assert rows[0]["Var_Lote"] == "L-9"
    # el booleano sin valor se escribe «No», no queda vacío
    assert rows[1]["Var_Marca propia"] == "No"
    # cantidad nula se convierte en cadena vacía
    assert rows[1]["cantidad"] == ""
    assert rows[2]["categoria"] == ""


def test_products_y_tasas_csv() -> None:
    assert products_rows(PRODUCTS)[0] == {
        "sku": "p1",
        "nombre": "Harina",
        "categoriaId": "c1",
        "marca": "Marca X",
        "unidad": "kg",
        "subcategoria": "Lácteos",
        "proveedor": "A",
    }
    # las tasas salen ordenadas por fecha
    assert [row["date"] for row in csv_rates_rows(RATES)] == ["2026-01-01", "2026-03-01"]


def test_nombres_de_fichero_csv() -> None:
    assert csv_filename("tx", NOW_MS) == f"transacciones_{NOW_DATE}.csv"
    assert csv_filename("prod", NOW_MS) == f"catalogo_{NOW_DATE}.csv"
    assert csv_filename("rates", NOW_MS) == f"tasas_{NOW_DATE}.csv"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def test_resumen_anual() -> None:
    rows = summary_rows(
        TXS,
        workspace="personal",
        year=2026,
        products=PRODUCTS,
        categories=CATEGORIES,
        rates=RATES,
        now_ms=NOW_MS,
    )
    values = {row["Campo"]: row["Valor"] for row in rows}
    assert values["Reporte"] == "Resumen Ejecutivo Anual"
    assert values["Entorno"] == "Personal"
    assert values["Generado"] == "15/3/2026, 12:30:45 p. m."
    assert values["Total transacciones"] == 3
    assert values["Total Bolívares (Bs)"] == 396
    assert values["Total Dólares (USD)"] == 11
    # 11 / 3 redondeado a 4 decimales
    assert values["Ticket promedio USD"] == 3.6667
    assert values["Base imponible acumulada Bs"] == 272.41
    assert values["IVA (16%) acumulado Bs"] == 43.59
    assert values["Gastos exentos Bs"] == 80
    assert values["Tasa BCV inicio año"] == 40
    assert values["Tasa BCV cierre año"] == 60
    assert values["Devaluación acumulada %"] == 50


def test_resumen_sin_tasas_del_año() -> None:
    rows = summary_rows(
        TXS,
        workspace="personal",
        year=2026,
        products=PRODUCTS,
        categories=CATEGORIES,
        rates=[],
        now_ms=NOW_MS,
    )
    values = {row["Campo"]: row["Valor"] for row in rows}
    assert values["Tasa BCV inicio año"] == "—"
    assert values["Devaluación acumulada %"] == 0


def test_libro_mayor_ordena_por_fecha_y_id() -> None:
    rows = ledger_rows(TXS, CATEGORIES)
    assert [row["ID"] for row in rows] == [1, 2, 3]
    first = rows[1]
    assert first["Categoría"] == "Comida"
    assert first["Cantidad"] == 2
    assert first["IVA Bs"] == 16
    assert first["Tasa Bs/USD"] == 29
    assert first["Var_Marca propia"] == "Sí"
    # la transacción sin base ni IVA se escribe con ceros redondeados
    assert rows[0]["Base Imponible Bs"] == 0
    assert rows[0]["IVA Bs"] == 0


def test_catalogo_usa_la_ultima_compra() -> None:
    rows = catalog_rows(PRODUCTS, TXS, CATEGORIES, now_ms=NOW_MS)
    first = rows[0]
    assert first["Última compra"] == "2026-03-05"
    assert first["Último precio Bs/unit"] == 58  # 116 / 2
    assert first["Último precio USD/unit"] == 2  # 4 / 2
    assert first["Compras (año)"] == 2
    assert first["Creado"] == "1/1/2026"
    # 5 / 3 redondeado a 4 decimales
    assert rows[1]["Última compra"] == "2026-03-06"
    assert rows[1]["Último precio USD/unit"] == 1.6667


def test_analisis_de_precios() -> None:
    rows = analysis_rows(TXS, PRODUCTS)
    by_sku = {row["SKU"]: row for row in rows}
    harina = by_sku["p1"]
    assert harina["Nº Compras"] == 2
    assert harina["Mín USD/unit"] == 2
    assert harina["Máx USD/unit"] == 2
    assert harina["Variación USD %"] == 0
    assert harina["CAGR USD %"] == 0
    assert harina["Primera compra"] == "2026-01-10"
    assert harina["Última compra"] == "2026-03-05"
    # el arroz (una sola compra) tiene dispersión cero
    assert by_sku["p2"]["Dispersión USD"] == 0


def test_hoja_de_categorias_y_pagos() -> None:
    rows = category_payment_rows(TXS, CATEGORIES)
    assert rows[0]["Sección"] == "CATEGORÍAS (con impacto inflacionario)"
    assert rows[1]["Sección"] == "Comida"
    assert rows[1]["Dato A"] == 2
    assert rows[1]["Dato C"] == 6
    assert rows[1]["Dato F"] == 0
    assert rows[2]["Sección"] == "Sin categoría"
    # 6 / 11 y 5 / 11
    assert rows[1]["Dato C"] == 6 and rows[2]["Dato C"] == 5
    assert any(row["Sección"] == "MÉTODOS DE PAGO" for row in rows)
    payments = rows[-1]
    assert payments["Sección"] == "No especificado"
    assert payments["Dato A"] == 3
    assert payments["Dato B"] == 396


def test_hoja_de_tasas_con_modelo() -> None:
    rows = rates_rows(RATES)
    assert rows[0] == {"Fecha": "2026-01-01", "Bs por USD": 40, "Fuente": "BCV"}
    assert rows[2] == {"Fecha": "", "Bs por USD": "", "Fuente": ""}
    assert rows[3]["Fecha"] == "MODELO REGRESIÓN"
    # y = 40 + (20/59)·días
    assert rows[4]["Bs por USD"] == "a=40.0000"
    assert rows[4]["Fuente"].startswith("b=0.3390 R²=1.000")
    assert rows[5]["Fecha"] == ""
    assert rows[6]["Fecha"] == "PROYECCIONES"
    assert rows[7]["Fecha"] == "Hoy + 7 días"
    # 40 + (20/59)·(59 + 7) = 62.3729
    assert rows[7]["Bs por USD"] == 62.3729
    assert rows[11]["Fecha"] == "Hoy + 365 días"


def test_tasas_sin_modelo_con_una_sola_tasa() -> None:
    rows = rates_rows([RATES[0]])
    assert len(rows) == 1


def test_anchos_de_columna() -> None:
    assert column_widths([]) == []
    assert column_widths([{"a": "x", "bb": "yyyy"}]) == [10, 10]
    assert column_widths([{"columna": "un valor bastante largo"}]) == [25]
    # el máximo es 42
    assert column_widths([{"c": "x" * 100}]) == [42]


def test_las_seis_hojas_y_el_marcador_de_analisis() -> None:
    sheets = build_sheets(
        TXS,
        workspace="personal",
        year=2026,
        products=PRODUCTS,
        categories=CATEGORIES,
        rates=RATES,
        now_ms=NOW_MS,
    )
    assert [name for name, _rows, _widths in sheets] == [
        "1. Resumen",
        "2. Libro Mayor",
        "3. Catálogo SKUs",
        "4. Análisis Precios",
        "5. Categorías y Pagos",
        "6. Tasas y Proyecciones",
    ]
    # la hoja 6 no pasa por autoWidth
    assert sheets[5][2] is None

    empty = build_sheets(
        [],
        workspace="personal",
        year=2026,
        products=[],
        categories=[],
        rates=[],
        now_ms=NOW_MS,
    )
    assert empty[3][1] == [{"Info": "Sin datos en el año"}]


def test_nombre_del_libro() -> None:
    assert xlsx_filename("personal", 2026, NOW_MS) == (
        f"Gastos_personal_2026_{NOW_DATE}.xlsx"
    )


def test_write_xlsx_escribe_un_libro_legible(tmp_path) -> None:
    sheets = build_sheets(
        TXS,
        workspace="personal",
        year=2026,
        products=PRODUCTS,
        categories=CATEGORIES,
        rates=RATES,
        now_ms=NOW_MS,
    )
    target = tmp_path / "libro.xlsx"
    write_xlsx(target, sheets)

    from openpyxl import load_workbook

    workbook = load_workbook(target, read_only=True)
    assert workbook.sheetnames == [name for name, _rows, _widths in sheets]
    sheet = workbook["2. Libro Mayor"]
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header[0] == "ID"
    assert "Tasa Bs/USD" in header
    workbook.close()


# ---------------------------------------------------------------------------
# Respaldo JSON
# ---------------------------------------------------------------------------


def test_respaldo_json() -> None:
    data = build_backup(
        transactions=TXS,
        products=PRODUCTS,
        categories=CATEGORIES,
        rates=RATES,
        closures=[],
        workspace="personal",
        fiscal_year=2026,
        ls_config={
            "theme": "dark",
            "workspace": "personal",
            "fiscalYear": "2026",
            "decimals": "2",
            "currency": "VES",
        },
        now_ms=NOW_MS,
    )
    assert data["meta"] == {
        "app": "GastosBimonetarios",
        "version": 1,
        "exportedAt": "2026-03-15T16:30:45.000Z",
        "workspace": "personal",
        "fiscalYear": 2026,
    }
    assert data["lsConfig"]["theme"] == "dark"
    assert data["lsConfig"]["currency"] == "VES"
    assert len(data["transactions"]) == 3
    assert backup_filename(NOW_MS) == f"backup_gastos_{NOW_DATE}.json"
    # el texto es JSON válido y redondo
    text = backup_text(data)
    assert json.loads(text)["meta"]["app"] == "GastosBimonetarios"
    assert text.startswith("{\n  \"meta\": {")


def test_backup_text_convierte_nan_en_null() -> None:
    assert backup_text({"a": float("nan"), "b": 1.0}) == (
        '{\n  "a": null,\n  "b": 1.0\n}'
    )


def test_validacion_del_respaldo() -> None:
    assert is_valid_backup({"meta": {"app": "GastosBimonetarios"}})
    assert not is_valid_backup({"meta": {"app": "otra"}})
    assert not is_valid_backup(None)


def test_importacion_json() -> None:
    data = {
        "meta": {"app": "GastosBimonetarios", "version": 1},
        "transactions": [{"id": 9, "amountBs": 100}, {"id": 10, "amountBs": 200}],
        "products": [{"id": "p1", "name": "Harina"}, {"id": "p9", "name": "Nuevo"}],
        "categories": [{"id": "c1", "name": "Comida"}, {"id": "c9", "name": "Nueva"}],
        "rates": [{"id": "r9", "date": "2026-01-01", "value": 41}],
        "closures": [{"id": "cl1", "year": 2025}],
    }
    plan = import_plan(
        data,
        existing_product_ids={"p1"},
        existing_category_ids={"c1"},
        existing_closure_ids={"cl1"},
    )
    # las transacciones pierden el id y se añaden todas
    assert plan["transactions"] == [{"amountBs": 100}, {"amountBs": 200}]
    assert plan["count"] == 2
    # lo existente no se toca; las tasas se escriben siempre
    assert [p["id"] for p in plan["products"]] == ["p9"]
    assert [c["id"] for c in plan["categories"]] == ["c9"]
    assert [c["id"] for c in plan["closures"]] == []
    assert [r["id"] for r in plan["rates"]] == ["r9"]

    assert import_prompt(data) == (
        "Importación de respaldo\n\n"
        "Transacciones: 2\nProductos: 2\nCategorías: 2\nTasas: 1\nCierres: 1\n\n"
        "¿Continuar?"
    )


def test_importacion_desde_excel() -> None:
    rows = [
        {
            "Fecha": "2026-03-05",
            "Producto": "Harina",
            "Total Bs": 116,
            "Total USD": 4,
            "Cantidad": 2,
            "Unidad": "kg",
            "Marca": "Marca X",
            "Modalidad IVA": "Exento",
            "Base Imponible Bs": 100,
            "IVA Bs": 16,
            "Tasa Bs/USD": 29,
        },
        {"Fecha": "2026-03-06", "Producto": "Arroz", "Total Bs": 200, "Cantidad": 3},
        {"Fecha": "", "Producto": "Sin fecha", "Total Bs": 100},
        {"Fecha": "2026-03-07", "Producto": "Sin monto", "Total Bs": 0},
        {"Fecha": "2026-03-08", "Producto": "harina", "Total Bs": 50},
    ]
    counter = {"n": 0}

    def next_sku() -> str:
        counter["n"] += 1
        return f"SKU-{counter['n']}"

    plan = plan_excel_import(
        rows,
        products=[],
        categories=CATEGORIES,
        workspace="personal",
        fiscal_year=2026,
        now_ms=NOW_MS,
        sku_factory=next_sku,
        app_rate_value=40,
    )
    # se descartan la fila sin fecha y la de importe cero
    assert plan["count"] == 3
    # «Harina» y «harina» son el mismo producto (búsqueda sin distinguir caja)
    assert [p["name"] for p in plan["products"]] == ["Harina", "Arroz"]
    assert [tx["productId"] for tx in plan["transactions"]] == ["SKU-1", "SKU-2", "SKU-1"]
    first = plan["transactions"][0]
    assert first["taxMode"] == "exento"
    assert first["categoryId"] == "c1"
    assert first["rate"] == 29
    assert first["amountUsd"] == 4
    assert first["unit"] == "kg"
    # sin «Tasa Bs/USD» usa la tasa activa
    assert plan["transactions"][1]["rate"] == 40
    assert plan["transactions"][1]["amountUsd"] == 5  # 200 / 40


def test_importacion_desde_excel_sin_tasa() -> None:
    plan = plan_excel_import(
        [{"Fecha": "2026-03-05", "Producto": "Harina", "Total Bs": 100}],
        products=[],
        categories=[],
        workspace="personal",
        fiscal_year=2026,
        now_ms=NOW_MS,
        sku_factory=lambda: "SKU-1",
    )
    # sin tasa activa la tasa es 1 y el total en Bs vale como USD
    assert plan["transactions"][0]["rate"] == 1
    assert plan["transactions"][0]["amountUsd"] == 100
    # sin categorías el producto se crea sin categoría
    assert plan["products"][0]["categoryId"] == ""


# ---------------------------------------------------------------------------
# Acta de cierre
# ---------------------------------------------------------------------------


def _snapshot(**overrides):
    base = {
        "id": "closure_1",
        "workspace": "personal",
        "year": 2026,
        "closedAt": NOW_MS,
        "totalBs": 1156,
        "totalUsd": 31.1659,
        "ivaTotal": 126.35,
        "baseTotal": 789.65,
        "txCount": 5,
        "rateStart": 40,
        "rateEnd": 60,
        "rateStartDate": "2026-01-01",
        "rateEndDate": "2026-03-01",
        "rateDelta": 50,
        "inflationAvg": 12.345,
        "topCategories": [
            {"name": "Comida & Bebida", "color": "#111111", "total": 17.1659, "count": 3}
        ],
        "topExpenses": [
            {"productName": 'Arroz "premium"', "date": "2026-03-06", "amountBs": 480, "amountUsd": 12}
        ],
        "notes": "Cierre <b>manual</b> & revisado",
    }
    base.update(overrides)
    return base


def test_acta_de_cierre() -> None:
    html = build_closure_act_html(_snapshot(), now_ms=NOW_MS)
    assert html.startswith("\n    <div class=\"acta-doc\">")
    assert html.endswith("</div>")
    assert "<h1>Acta de Cierre de Ejercicio Fiscal 2026</h1>" in html
    assert "<b>Entorno:</b> Personal " in html
    assert "<b>Fecha de cierre:</b> 15/3/2026, 12:30:45 p. m.<br>" in html
    # los totales de Bs llevan separador de miles; los USD usan toFixed
    assert '<span class="val">Bs. 1.156,00</span>' in html
    assert '<span class="val">$31.17</span>' in html
    assert "<span class=\"lbl\">Inflación micro prom.</span><span class=\"val\">12.35%</span>" in html
    assert "<div class=\"acta-kpi\"><span class=\"lbl\">Base imponible</span>" in html
    # datos escapados
    assert "Comida &amp; Bebida" in html
    assert "Arroz &quot;premium&quot;" in html
    assert "Cierre &lt;b&gt;manual&lt;/b&gt; &amp; revisado" in html
    # tasas y devaluación
    assert ">40,00<" in html
    assert ">60,00<" in html
    assert "<b>50.00%</b>" in html
    assert "Emitido: 15/3/2026, 12:30:45 p. m." in html


def test_acta_sin_datos() -> None:
    html = build_closure_act_html(
        _snapshot(
            workspace="trabajo",
            year=2025,
            closedAt="2025-12-31T23:30:00",
            totalBs=0,
            totalUsd=0,
            ivaTotal=0,
            baseTotal=0,
            txCount=0,
            rateStart=None,
            rateEnd=None,
            rateStartDate=None,
            rateEndDate=None,
            rateDelta=0,
            inflationAvg=0,
            topCategories=[],
            topExpenses=[],
            notes="",
        ),
        now_ms=NOW_MS,
    )
    assert "Acta de Cierre de Ejercicio Fiscal 2025" in html
    assert "<b>Entorno:</b> Profesional " in html
    assert "<p><i>Sin categorías registradas.</i></p>" in html
    assert "<p><i>Sin gastos registrados.</i></p>" in html
    assert "<i>Sin observaciones registradas.</i>" in html
    # sin tasas se escribe un guion largo
    assert html.count('<td class="num">—</td>') == 2
    assert ">Bs. 0,00<" in html


@pytest.mark.parametrize(
    "closed_at, expected",
    [
        ("2025-12-31T23:30:00", "31/12/2025, 11:30:00 p. m."),
        ("2026-01-01T00:00:00", "1/1/2026, 12:00:00 a. m."),
    ],
)
def test_acta_fecha_de_cierre(closed_at: str, expected: str) -> None:
    html = build_closure_act_html(_snapshot(closedAt=closed_at), now_ms=NOW_MS)
    assert f"<b>Fecha de cierre:</b> {expected}<br>" in html
