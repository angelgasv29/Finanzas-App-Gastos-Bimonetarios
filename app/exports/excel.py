"""Exportación Excel de 6 hojas (``exportXLSX`` y ``autoWidth``, 6329-6573).

La versión de referencia usa SheetJS en el navegador; esta implementación usa ``openpyxl``. Lo que se
verifica no es el fichero binario sino **el contenido de las hojas**: cada hoja
se construye como una lista de diccionarios (clave = cabecera) y el oráculo
ejecuta ``exportXLSX`` con un SheetJS falso que graba exactamente esas listas.
"""

from __future__ import annotations

from typing import Any

from ..core.js_locale import locale_date, locale_datetime
from ..core.money import js_number_or_zero, js_round, js_to_fixed
from ..core.periods import date_to_iso
from ..core.stats import cagr, linear_regression, mean, median, stdev, unit_price
from ..domain.dates import local_date_at, utc_midnight_ms
from .csv_export import transactions_rows  # noqa: F401  (reexportado por comodidad)

__all__ = [
    "SHEET_NAMES",
    "summary_rows",
    "ledger_rows",
    "catalog_rows",
    "analysis_rows",
    "category_payment_rows",
    "rates_rows",
    "column_widths",
    "build_sheets",
    "xlsx_filename",
    "write_xlsx",
]

MS_PER_DAY = 86_400_000
PROJECTION_HORIZONS = (7, 30, 90, 180, 365)


def _has_custom_vat(txs: list[dict[str, Any]]) -> bool:
    """¿Alguna compra usa una alícuota distinta de 16 %?"""
    for tx in txs:
        stored = tx.get("ivaRate")
        if stored is not None:
            try:
                rate = float(stored)
            except (TypeError, ValueError):
                continue
        else:
            base = js_number_or_zero(tx.get("baseImponible"))
            iva = js_number_or_zero(tx.get("ivaMonto"))
            rate = iva / base if base > 0 else None
        if rate is not None and round(rate * 100, 2) != 16.0:
            return True
    return False


def _vat_percent(tx: dict[str, Any]) -> Any:
    """Alícuota en porcentaje; se deduce de base/IVA si es un registro viejo."""
    stored = tx.get("ivaRate")
    if stored is not None:
        try:
            return js_round(float(stored) * 100, 2)
        except (TypeError, ValueError):
            return ""
    base = js_number_or_zero(tx.get("baseImponible"))
    iva = js_number_or_zero(tx.get("ivaMonto"))
    if base > 0:
        return js_round((iva / base) * 100, 2)
    return ""

SHEET_NAMES = (
    "1. Resumen",
    "2. Libro Mayor",
    "3. Catálogo SKUs",
    "4. Análisis Precios",
    "5. Categorías y Pagos",
    "6. Tasas y Proyecciones",
)


def _by_id(items: list[dict[str, Any]]) -> dict[Any, dict[str, Any]]:
    return {item.get("id"): item for item in items}


def _days_between(start: Any, end: Any) -> float:
    """``(new Date(end) - new Date(start)) / 86400000`` (medianoche UTC)."""
    first = utc_midnight_ms(start)
    second = utc_midnight_ms(end)
    if first is None or second is None:
        return 0.0
    return (second - first) / MS_PER_DAY


def _sort_year_txs(year_txs: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """``sort((a,b) => a.date.localeCompare(b.date) || (a.id - b.id))``."""
    return sorted(
        year_txs,
        key=lambda tx: (tx.get("date") or "", js_number_or_zero(tx.get("id"))),
    )


def summary_rows(
    year_txs: list[dict[str, Any]],
    *,
    workspace: str,
    year: int,
    products: list[dict[str, Any]],
    categories: list[dict[str, Any]],
    rates: list[dict[str, Any]],
    now_ms: float,
    locale: str | None = None,
) -> list[dict[str, Any]]:
    """Hoja «1. Resumen»: totales del año, IVA y devaluación de la tasa."""
    total_bs = 0.0
    total_usd = 0.0
    iva_total = 0.0
    base_total = 0.0
    exento_total = 0.0
    for tx in year_txs:
        total_bs += js_number_or_zero(tx.get("amountBs"))
        total_usd += js_number_or_zero(tx.get("amountUsd"))
        iva_total += js_number_or_zero(tx.get("ivaMonto"))
        base_total += js_number_or_zero(tx.get("baseImponible"))
        if tx.get("taxMode") == "exento":
            exento_total += js_number_or_zero(tx.get("amountBs"))

    year_rates = [
        rate
        for rate in rates
        if str(rate.get("date") or "").startswith(str(year))
    ]
    sorted_rates = sorted(year_rates, key=lambda rate: rate.get("date") or "")
    rate_first = js_number_or_zero(sorted_rates[0].get("value")) if sorted_rates else 0.0
    rate_last = (
        js_number_or_zero(sorted_rates[-1].get("value")) if sorted_rates else 0.0
    )
    rate_delta = ((rate_last - rate_first) / rate_first) * 100 if rate_first else 0.0
    iva_label = (
        "IVA acumulado Bs"
        if _has_custom_vat(year_txs)
        else "IVA (16%) acumulado Bs"
    )

    return [
        {"Campo": "Reporte", "Valor": "Resumen Ejecutivo Anual"},
        {
            "Campo": "Entorno",
            "Valor": "Personal" if workspace == "personal" else "Profesional",
        },
        {"Campo": "Año fiscal", "Valor": year},
        {"Campo": "Generado", "Valor": locale_datetime(now_ms, locale)},
        {"Campo": "", "Valor": ""},
        {"Campo": "Total transacciones", "Valor": len(year_txs)},
        {"Campo": "Total Bolívares (Bs)", "Valor": js_round(total_bs, 2)},
        {"Campo": "Total Dólares (USD)", "Valor": js_round(total_usd, 4)},
        {
            "Campo": "Ticket promedio USD",
            "Valor": js_round(total_usd / len(year_txs), 4) if year_txs else 0,
        },
        {"Campo": "", "Valor": ""},
        {"Campo": "Base imponible acumulada Bs", "Valor": js_round(base_total, 2)},
        {"Campo": iva_label, "Valor": js_round(iva_total, 2)},
        {"Campo": "Gastos exentos Bs", "Valor": js_round(exento_total, 2)},
        {"Campo": "", "Valor": ""},
        {"Campo": "Tasa BCV inicio año", "Valor": rate_first or "—"},
        {"Campo": "Tasa BCV cierre año", "Valor": rate_last or "—"},
        {"Campo": "Devaluación acumulada %", "Valor": js_round(rate_delta, 2)},
        {"Campo": "", "Valor": ""},
        {"Campo": "Productos en catálogo", "Valor": len(products)},
        {"Campo": "Categorías activas", "Valor": len(categories)},
    ]


def ledger_rows(
    year_txs: list[dict[str, Any]], categories: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Hoja «2. Libro Mayor»: una fila por transacción del año."""
    category_map = _by_id(categories)
    show_vat = _has_custom_vat(year_txs)
    rows: list[dict[str, Any]] = []
    for tx in _sort_year_txs(year_txs):
        category = category_map.get(tx.get("categoryId"))
        row: dict[str, Any] = {
            "ID": tx.get("id"),
            "Fecha": tx.get("date"),
            "Producto": tx.get("productName"),
            "SKU": tx.get("productId") or "",
            "Categoría": category.get("name") if category else "",
            "Subcategoría": tx.get("subcategory") or "",
            "Marca": tx.get("brand") or "",
            "Comercio": tx.get("provider") or "",
            "Método de Pago": tx.get("paymentMethod") or "",
            "Plataforma Origen": tx.get("originPlatform") or "",
            "Naturaleza": tx.get("nature") or "",
            "Cantidad": tx.get("quantity") if tx.get("quantity") is not None else "",
            "Unidad": tx.get("unit") or "",
            "Modalidad IVA": tx.get("taxMode"),
            "Base Imponible Bs": js_round(tx.get("baseImponible") or 0, 2),
            "IVA Bs": js_round(tx.get("ivaMonto") or 0, 2),
            "Total Bs": js_round(tx.get("amountBs"), 2),
            "Total USD": js_round(tx.get("amountUsd"), 4),
            "Tasa Bs/USD": js_round(tx.get("rate") or 0, 4),
            "Fecha Tasa": tx.get("rateDate") or "",
            "Fuente Tasa": tx.get("rateSource") or "",
            "Notas": tx.get("notes") or "",
        }
        if show_vat:
            with_vat: dict[str, Any] = {}
            for key, value in row.items():
                with_vat[key] = value
                if key == "Modalidad IVA":
                    with_vat["Tasa IVA %"] = _vat_percent(tx)
            row = with_vat
        for field in (category or {}).get("customFields") or []:
            values = tx.get("customValues") or {}
            value = values.get(field.get("id"), "")
            row[f"Var_{field.get('label')}"] = (
                ("Sí" if value else "No") if field.get("type") == "boolean" else value
            )
        rows.append(row)
    return rows


def catalog_rows(
    products: list[dict[str, Any]],
    year_txs: list[dict[str, Any]],
    categories: list[dict[str, Any]],
    *,
    now_ms: float,
    locale: str | None = None,
) -> list[dict[str, Any]]:
    """Hoja «3. Catálogo SKUs»: ficha de cada producto con su última compra."""
    category_map = _by_id(categories)
    rows: list[dict[str, Any]] = []
    for product in products:
        product_txs = [
            tx for tx in year_txs if tx.get("productId") == product.get("id")
        ]
        ordered = sorted(
            product_txs, key=lambda tx: tx.get("date") or "", reverse=True
        )
        last_tx = ordered[0] if ordered else None
        category = category_map.get(product.get("categoryId"))
        last_unit = unit_price(last_tx) if last_tx else None
        rows.append(
            {
                "SKU": product.get("id"),
                "Nombre": product.get("name"),
                "Categoría": category.get("name") if category else "",
                "Marca": product.get("brand") or "",
                "Unidad": product.get("unit") or "",
                "Subcategoría": product.get("subcategory") or "",
                "Proveedor habitual": product.get("provider") or "",
                "Compras (año)": len(product_txs),
                "Último precio Bs/unit": (
                    js_round(last_unit["unitBs"], 2) if last_unit else ""
                ),
                "Último precio USD/unit": (
                    js_round(last_unit["unitUsd"], 4) if last_unit else ""
                ),
                "Última compra": last_tx.get("date") if last_tx else "",
                "Creado": locale_date(
                    js_number_or_zero(product.get("createdAt")), locale
                ),
            }
        )
    return rows


def analysis_rows(
    year_txs: list[dict[str, Any]], products: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Hoja «4. Análisis Precios»: estadística de precio por SKU."""
    product_map = _by_id(products)
    grouped: dict[Any, list[dict[str, Any]]] = {}
    for tx in year_txs:
        product_id = tx.get("productId")
        if not product_id:
            continue
        grouped.setdefault(product_id, []).append(tx)

    rows: list[dict[str, Any]] = []
    for product_id, purchases in grouped.items():
        units_usd = [unit_price(tx)["unitUsd"] for tx in purchases]
        units_bs = [unit_price(tx)["unitBs"] for tx in purchases]
        ordered = sorted(purchases, key=lambda tx: tx.get("date") or "")
        first_unit = unit_price(ordered[0])
        last_unit = unit_price(ordered[-1])
        delta_usd = (
            ((last_unit["unitUsd"] - first_unit["unitUsd"]) / first_unit["unitUsd"]) * 100
            if first_unit["unitUsd"]
            else 0.0
        )
        delta_bs = (
            ((last_unit["unitBs"] - first_unit["unitBs"]) / first_unit["unitBs"]) * 100
            if first_unit["unitBs"]
            else 0.0
        )
        product = product_map.get(product_id)
        days = _days_between(ordered[0].get("date"), ordered[-1].get("date"))
        rows.append(
            {
                "SKU": product_id,
                "Producto": product.get("name") if product else product_id,
                "Nº Compras": len(purchases),
                "Mín USD/unit": js_round(min(units_usd), 4),
                "Máx USD/unit": js_round(max(units_usd), 4),
                "Prom USD/unit": js_round(mean(units_usd), 4),
                "Mediana USD/unit": js_round(median(units_usd), 4),
                "Dispersión USD": js_round(stdev(units_usd), 4),
                "Variación USD %": js_round(delta_usd, 2),
                "Variación Bs %": js_round(delta_bs, 2),
                "CAGR USD %": js_round(
                    cagr(first_unit["unitUsd"], last_unit["unitUsd"], days), 2
                ),
                "Primera compra": ordered[0].get("date"),
                "Última compra": ordered[-1].get("date"),
            }
        )
    rows.sort(key=lambda row: row["Variación USD %"], reverse=True)
    return rows


def category_payment_rows(
    year_txs: list[dict[str, Any]], categories: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    """Hoja «5. Categorías y Pagos», con impacto inflacionario por categoría."""
    category_map = _by_id(categories)

    by_category: dict[str, dict[str, Any]] = {}
    for tx in year_txs:
        category = category_map.get(tx.get("categoryId"))
        key = category.get("name") if category else "Sin categoría"
        group = by_category.get(key) if key else None
        if group is None:
            group = {
                "name": key,
                "count": 0,
                "totalUsd": 0.0,
                "totalBs": 0.0,
                "iva": 0.0,
                "base": 0.0,
                "records": [],
            }
            by_category[key] = group
        group["count"] += 1
        group["totalUsd"] += js_number_or_zero(tx.get("amountUsd"))
        group["totalBs"] += js_number_or_zero(tx.get("amountBs"))
        group["iva"] += js_number_or_zero(tx.get("ivaMonto"))
        group["base"] += js_number_or_zero(tx.get("baseImponible"))
        group["records"].append(tx)

    total_usd = 0.0
    for group in by_category.values():
        total_usd += group["totalUsd"]

    ordered_categories = sorted(
        by_category.values(), key=lambda group: group["totalUsd"], reverse=True
    )

    category_sheet: list[dict[str, Any]] = []
    for group in ordered_categories:
        sku_in_category: dict[Any, list[dict[str, Any]]] = {}
        for record in group["records"]:
            product_id = record.get("productId")
            bucket = sku_in_category.get(product_id) if product_id else None
            if bucket is None:
                bucket = []
                sku_in_category[product_id] = bucket
            bucket.append(record)

        extra_inflation = 0.0
        for purchases in sku_in_category.values():
            if len(purchases) < 2:
                continue
            purchases.sort(key=lambda tx: tx.get("date") or "")
            first_unit = unit_price(purchases[0])["unitUsd"]
            last = purchases[-1]
            last_unit = unit_price(last)["unitUsd"]
            last_quantity = js_number_or_zero(last.get("quantity")) or 1
            extra_inflation += (last_unit - first_unit) * last_quantity

        category_sheet.append(
            {
                "Categoría": group["name"],
                "Compras": group["count"],
                "Total Bs": js_round(group["totalBs"], 2),
                "Total USD": js_round(group["totalUsd"], 4),
                "Base Imponible Bs": js_round(group["base"], 2),
                "IVA Bs": js_round(group["iva"], 2),
                "% del gasto USD": (
                    js_round((group["totalUsd"] / total_usd) * 100, 2)
                    if total_usd
                    else 0
                ),
                "Impacto Inflacionario +$ USD": js_round(extra_inflation, 2),
            }
        )

    by_payment: dict[str, dict[str, Any]] = {}
    for tx in year_txs:
        key = tx.get("paymentMethod") or "No especificado"
        group = by_payment.get(key)
        if group is None:
            group = {"count": 0, "totalUsd": 0.0, "totalBs": 0.0}
            by_payment[key] = group
        group["count"] += 1
        group["totalUsd"] += js_number_or_zero(tx.get("amountUsd"))
        group["totalBs"] += js_number_or_zero(tx.get("amountBs"))

    payment_sheet = sorted(
        by_payment.items(), key=lambda item: item[1]["totalUsd"], reverse=True
    )

    empty_row = {
        "Sección": "",
        "Dato A": "",
        "Dato B": "",
        "Dato C": "",
        "Dato D": "",
        "Dato E": "",
        "Dato F": "",
    }
    rows: list[dict[str, Any]] = [
        {
            "Sección": "CATEGORÍAS (con impacto inflacionario)",
            "Dato A": "",
            "Dato B": "",
            "Dato C": "",
            "Dato D": "",
            "Dato E": "",
            "Dato F": "",
        }
    ]
    rows.extend(
        {
            "Sección": category["Categoría"],
            "Dato A": category["Compras"],
            "Dato B": category["Total Bs"],
            "Dato C": category["Total USD"],
            "Dato D": category["Base Imponible Bs"],
            "Dato E": category["IVA Bs"],
            "Dato F": category["Impacto Inflacionario +$ USD"],
        }
        for category in category_sheet
    )
    rows.append(dict(empty_row))
    rows.append(
        {
            "Sección": "MÉTODOS DE PAGO",
            "Dato A": "",
            "Dato B": "",
            "Dato C": "",
            "Dato D": "",
            "Dato E": "",
            "Dato F": "",
        }
    )
    rows.extend(
        {
            "Sección": key,
            "Dato A": group["count"],
            "Dato B": js_round(group["totalBs"], 2),
            "Dato C": js_round(group["totalUsd"], 4),
            "Dato D": "",
            "Dato E": "",
            "Dato F": "",
        }
        for key, group in payment_sheet
    )
    return rows


def rates_rows(rates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Hoja «6. Tasas y Proyecciones»: histórico, modelo y proyección."""
    sorted_rates = sorted(rates, key=lambda rate: rate.get("date") or "")
    rows: list[dict[str, Any]] = [
        {
            "Fecha": rate.get("date"),
            "Bs por USD": js_round(rate.get("value"), 4),
            "Fuente": rate.get("source") or "",
        }
        for rate in sorted_rates
    ]
    if len(sorted_rates) >= 2:
        base_date = sorted_rates[0].get("date")
        xs = [_days_between(base_date, rate.get("date")) for rate in sorted_rates]
        ys = [js_number_or_zero(rate.get("value")) for rate in sorted_rates]
        regression = linear_regression(xs, ys)
        last_x = xs[-1]
        rows.append({"Fecha": "", "Bs por USD": "", "Fuente": ""})
        rows.append({"Fecha": "MODELO REGRESIÓN", "Bs por USD": "", "Fuente": ""})
        rows.append(
            {
                "Fecha": "y = a + b·días",
                "Bs por USD": f"a={js_to_fixed(regression['a'], 4)}",
                "Fuente": (
                    f"b={js_to_fixed(regression['b'], 4)} "
                    f"R²={js_to_fixed(regression['r2'], 3)}"
                ),
            }
        )
        rows.append({"Fecha": "", "Bs por USD": "", "Fuente": ""})
        rows.append({"Fecha": "PROYECCIONES", "Bs por USD": "", "Fuente": ""})
        for horizon in PROJECTION_HORIZONS:
            projected = regression["a"] + regression["b"] * (last_x + horizon)
            rows.append(
                {
                    "Fecha": f"Hoy + {horizon} días",
                    "Bs por USD": js_round(projected, 4),
                    "Fuente": f"Proyección (R²={js_to_fixed(regression['r2'], 3)})",
                }
            )
    return rows


def column_widths(rows: list[dict[str, Any]]) -> list[int]:
    """``autoWidth``: ancho por columna, entre 10 y 42.

    >>> column_widths([{"a": "x", "bb": "yyyy"}])
    [10, 10]
    >>> column_widths([{"columna": "un valor bastante largo"}])
    [25]
    """
    if not rows:
        return []
    columns = list(rows[0].keys())
    widths: list[int] = []
    for column in columns:
        longest = len(str(column))
        for row in rows:
            value = row.get(column)
            longest = max(longest, len("" if value is None else str(value)))
        widths.append(min(42, max(10, longest + 2)))
    return widths


def build_sheets(
    year_txs: list[dict[str, Any]],
    *,
    workspace: str,
    year: int,
    products: list[dict[str, Any]],
    categories: list[dict[str, Any]],
    rates: list[dict[str, Any]],
    now_ms: float,
    locale: str | None = None,
) -> list[tuple[str, list[dict[str, Any]], list[int] | None]]:
    """Las 6 hojas en orden: nombre, filas y anchos (``None`` sin ``autoWidth``)."""
    summary = summary_rows(
        year_txs,
        workspace=workspace,
        year=year,
        products=products,
        categories=categories,
        rates=rates,
        now_ms=now_ms,
        locale=locale,
    )
    ledger = ledger_rows(year_txs, categories)
    catalog = catalog_rows(
        products, year_txs, categories, now_ms=now_ms, locale=locale
    )
    analysis = analysis_rows(year_txs, products)
    categories_sheet = category_payment_rows(year_txs, categories)
    rate_sheet = rates_rows(rates)

    analysis_sheet = analysis if analysis else [{"Info": "Sin datos en el año"}]
    return [
        (SHEET_NAMES[0], summary, column_widths(summary)),
        (SHEET_NAMES[1], ledger, column_widths(ledger)),
        (SHEET_NAMES[2], catalog, column_widths(catalog)),
        (SHEET_NAMES[3], analysis_sheet, column_widths(analysis)),
        (SHEET_NAMES[4], categories_sheet, column_widths(categories_sheet)),
        (SHEET_NAMES[5], rate_sheet, None),
    ]


def xlsx_filename(workspace: str, year: int, now_ms: float) -> str:
    """``Gastos_<entorno>_<año>_<YYYY-MM-DD>.xlsx``.

    >>> xlsx_filename("personal", 2026, 1773577845000)
    'Gastos_personal_2026_2026-03-15.xlsx'
    """
    return f"Gastos_{workspace}_{year}_{date_to_iso(local_date_at(now_ms))}.xlsx"


def write_xlsx(
    path: Any,
    sheets: list[tuple[str, list[dict[str, Any]], list[int] | None]],
) -> Any:
    """Escribe el libro con ``openpyxl`` respetando los anchos de ``autoWidth``."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows, widths in sheets:
        sheet = workbook.create_sheet(name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        if widths:
            for index, width in enumerate(widths, start=1):
                sheet.column_dimensions[get_column_letter(index)].width = width
    if not workbook.sheetnames:
        workbook.create_sheet("1. Resumen")
    workbook.save(str(path))
    return workbook
